"""Bounded spreadsheet parsing; preview never writes student records."""
import csv
import io
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile
from rest_framework.exceptions import ValidationError

FIELDS = {"lrn":"lrn", "first name":"firstName", "middle name":"middleName", "last name":"lastName", "sex":"sex", "birthday":"birthday", "grade":"grade", "section":"section", "status":"status", "school year":"schoolYear", "guardian":"guardian", "contact":"contact", "address":"address", "available credentials":"availableCredentials"}

def parse_students(file):
    if not file or not 0 < file.size <= 10*1024*1024:
        raise ValidationError("Choose a non-empty Excel or CSV file up to 10 MB.")
    suffix = Path(file.name).suffix.lower()
    book = None
    try:
        if suffix == ".xlsx":
            with ZipFile(file) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50*1024*1024:
                    raise ValidationError("Expanded workbook is too large.")
            file.seek(0)
            from openpyxl import load_workbook
            book = load_workbook(file, read_only=True, data_only=False)
            sheet = book.worksheets[0]
            def cells():
                for row in sheet.iter_rows(max_row=502, max_col=50):
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError("Replace formulas with values before importing.")
                    yield [cell.value for cell in row]
            source = cells()
        elif suffix == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=file.read(), on_demand=True)
            sheet = book.sheet_by_index(0)
            def cells():
                for index in range(min(sheet.nrows, 502)):
                    yield [xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value for cell in sheet.row(index)]
            source = cells()
        elif suffix == ".csv":
            source = csv.reader(io.StringIO(file.read().decode("utf-8-sig")), strict=True)
        else:
            raise ValidationError("Choose an .xlsx, .xls or UTF-8 .csv file.")
        def text(value):
            if value is None: return ""
            if isinstance(value, (datetime, date)): return value.strftime("%Y-%m-%d")
            if isinstance(value, float) and value.is_integer(): return str(int(value))
            return str(value).strip()
        headers = [text(v).lower().replace("_", " ") for v in next(source, [])]
        for required in ["lrn", "first name", "last name"]:
            if required not in headers: raise ValidationError(f"Missing column: {required}. Use the import template.")
        if len([h for h in headers if h in FIELDS]) != len(set(h for h in headers if h in FIELDS)):
            raise ValidationError("Duplicate column headings are not allowed.")
        rows, seen = [], set()
        for number, values in enumerate(source, 2):
            if number > 502: raise ValidationError("Import at most 500 students at a time.")
            if not any(text(v) for v in values): continue
            row = {FIELDS[h]: text(values[i]) if i < len(values) else "" for i,h in enumerate(headers) if h in FIELDS}
            lrn = row.get("lrn", "")
            if len(lrn) != 12 or not lrn.isascii() or not lrn.isdigit(): raise ValidationError(f"Row {number}: LRN must be exactly 12 digits. Format the Excel LRN column as Text to preserve leading zeros.")
            if not row.get("firstName") or not row.get("lastName"): raise ValidationError(f"Row {number}: first and last names are required.")
            if lrn in seen: raise ValidationError(f"Row {number}: duplicate LRN {lrn} in this file.")
            seen.add(lrn)
            row["availableCredentials"] = [v.strip() for v in row.get("availableCredentials", "").split(";") if v.strip()]
            row["status"] = row.get("status") or "Active"
            if row["status"] not in {"Active", "Graduated", "Archived", "Transferred"}: raise ValidationError(f"Row {number}: invalid status.")
            row.update(student=" ".join(row.get(k, "") for k in ["firstName", "middleName", "lastName"]).strip(), gradeSection=" - ".join(filter(None,[row.get("grade"),row.get("section")])), credential=", ".join(row["availableCredentials"]))
            rows.append(row)
        if not 1 <= len(rows) <= 500: raise ValidationError("Import 1 to 500 students at a time.")
        return rows
    except ValidationError: raise
    except Exception: raise ValidationError("Could not read this spreadsheet. Use the template and save an unencrypted Excel or UTF-8 CSV file.")
    finally:
        if book:
            if hasattr(book, "close"): book.close()
            elif hasattr(book, "release_resources"): book.release_resources()
